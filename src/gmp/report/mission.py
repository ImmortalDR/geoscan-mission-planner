"""Per-result PDF/DOCX export with explicit evidence and model limitations."""

from __future__ import annotations

import io
import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


LIMITATIONS = (
    "Проверка относится к заявленной расчётной модели, а не к разрешению на реальный полёт.",
    "Траектории линейны между точками в XYZ и времени; высота поверхности постоянна внутри ячейки растра.",
    "Динамика аппарата, перекрытие отдельных фотоснимков и качество бокового перекрытия не сертифицированы.",
    "Полоса съёмки вычисляется по высоте под маршрутом; затенение рельефом вне линии полёта не проверено.",
    "Глобальная оптимальность не доказана. Эталон и новый результат могут иметь разные допустимые маршруты.",
    "DSM включает здания и растительность и не является точным актуальным рельефом голой земли.",
    "Время завершения включает заданные подготовку и финальное снятие данных. Отсутствующие во входе значения считаются нулевыми; подготовка группы моделируется параллельно.",
)


def _value(value):
    if value is None:
        return "Не указано"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, float):
        return f"{value:,.3f}".replace(",", " ")
    return str(value)


def _sections(result: dict, validation: dict, metadata: dict) -> list[tuple[str, list]]:
    metrics = validation.get("metrics", {})
    certificate = validation.get("certificate")
    result_digest = hashlib.sha256(json.dumps(result, sort_keys=True, separators=(",", ":"),
                                             ensure_ascii=False, allow_nan=False).encode()).hexdigest()
    bound = bool(certificate) and certificate.get("result_sha256") == result_digest and certificate.get("input_sha256") == validation.get("input_sha256")
    safe = validation.get("passed") is True and validation.get("status") == "SAFE" and bound
    status = "SAFE" if safe else ("INFEASIBLE" if validation.get("passed") is True and validation.get("status") == "INFEASIBLE" else "UNSAFE")
    mode = metadata.get("mode", result.get("mode", "Не указан"))
    mode = {"live": "Реальный расчёт", "reference": "Готовый эталон", "fixture": "Готовый эталон"}.get(mode, mode)
    sections = [("Паспорт расчёта", [
        ("Сцена", result.get("scene_id")), ("Название", metadata.get("description", metadata.get("name"))),
        ("Целевая функция", {"makespan": "Время завершения миссии", "total_flight": "Суммарное полётное время"}.get(result.get("objective"), result.get("objective"))),
        ("Режим", mode), ("Проверенный статус", status),
        ("Заявленный статус", result.get("status")), ("Сформирован", datetime.now(timezone.utc).isoformat()),
        ("Валидатор", validation.get("validator_version")),
        ("Вход SHA-256", validation.get("input_sha256")), ("Результат SHA-256", validation.get("result_sha256")),
    ]), ("Проверенные показатели", [
        ("Покрытие исходных областей, %", metrics.get("coverage_percent")),
        ("Время завершения работ от начала окна, с", metrics.get("makespan_s")),
        ("Подготовка до первого вылета, с", metrics.get("preparation_time_s", 0)),
        ("Снятие данных после последней посадки, с", metrics.get("data_download_time_s", 0)),
        ("Обе длительности явно заданы во входе", "Да" if metrics.get("ground_operations_explicit") else "Нет: отсутствующие значения приняты равными 0"),
        ("Суммарное полётное время, с", metrics.get("total_flight_s")),
        ("Горизонтальная длина маршрутов, м", metrics.get("distance_m")),
        ("Количество вылетов", metrics.get("sortie_count")),
        ("Требуемая площадь, м²", metrics.get("required_area_m2")),
        ("Достижимость посадки", {"proved_by_resource_feasible_verified_suffix": "Подтверждена проверенным остатком маршрута", "not_proved": "Не подтверждена"}.get(metrics.get("landing_reachability"), metrics.get("landing_reachability"))),
    ])]
    schedule = []
    for s in result.get("sorties", []):
        schedule.append((s.get("id", "Вылет"),
                         f"БВС: {s.get('uav_id')}; {s.get('start_site')} -> {s.get('landing_site')}; "
                         f"{s.get('t_start')} / {s.get('t_end')}; время {_value(s.get('flight_time_s'))} с"))
    sections.append(("Расписание", schedule or [("Вылеты", "Исполняемый маршрут отсутствует")]))
    findings = [(v.get("code", "Нарушение"), v.get("message", v)) for v in validation.get("violations", [])]
    if certificate and not bound:
        findings.append(("CERTIFICATE_MISMATCH", "Сертификат не соответствует этому результату или входу"))
    if result.get("diagnosis"):
        findings.append(("Обоснование", result["diagnosis"]))
    sections.append(("Проверки и обоснование", findings or [("Нарушения", "Не обнаружены в пределах указанной модели")]))
    if certificate and safe:
        sections.append(("Расчётный Safety Certificate", [
            ("Идентификатор", certificate.get("fingerprint_sha256")),
            ("Выдан", certificate.get("issued_at")),
            ("Проверки", certificate.get("checks")),
            ("Допуск покрытия", certificate.get("coverage_policy")),
            ("Версия исходного кода проверки", certificate.get("validator_source_sha256")),
        ]))
    else:
        sections.append(("Расчётный Safety Certificate", [("Выдача", "Не выдан")]))
    provenance = [(k, metadata[k]) for k in ("metric_crs", "height_reference", "surface_model",
                  "real_elevation", "synthetic_mission", "not_operational_data", "user_assumptions",
                  "recommendation_assumptions", "recommendation_changes") if k in metadata]
    terrain = metadata.get("terrain", {})
    provenance.extend((key, terrain[key]) for key in ("source_id", "source_url", "source_sha256", "sha256",
                      "source_urls", "attribution", "license_notice", "license_url", "resolution_m", "vertical_datum") if key in terrain)
    provenance.extend((f"Источник DSM {index + 1}", source) for index, source in enumerate(terrain.get("sources", [])))
    sections.append(("Источники и допущения", provenance or [("Источники", "Не указаны во входных метаданных")]))
    sections.append(("Границы применимости", [(str(i), text) for i, text in enumerate(LIMITATIONS, 1)]))
    return sections


def _font() -> str:
    name = "H3DejaVuSans"
    if name not in pdfmetrics.getRegisteredFontNames():
        path = Path(os.environ.get("GMP_REPORT_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
        if not path.is_file():
            raise RuntimeError("A Cyrillic-capable font is required: install fonts-dejavu-core or set GMP_REPORT_FONT")
        pdfmetrics.registerFont(TTFont(name, str(path)))
    return name


def generate_report_bytes(format: str, result: dict, validation: dict, scene_metadata: dict) -> bytes:
    """Render a specific calculation; no historical test summary is substituted."""
    sections = _sections(result, validation, scene_metadata)
    output = io.BytesIO()
    if format.lower() == "docx":
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "DejaVu Sans"
        normal.font.size = Pt(10)
        for section in document.sections:
            section.top_margin = section.bottom_margin = Cm(1.8)
            section.left_margin = section.right_margin = Cm(2)
        document.add_heading("План полётного задания", 0)
        document.add_paragraph("Geoscan Mission Planner · расчётный отчёт H3")
        for title, rows in sections:
            document.add_heading(title, 1)
            table = document.add_table(rows=0, cols=2)
            table.style = "Light Shading Accent 1"
            for key, value in rows:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = _value(key), _value(value)
        document.save(output)
    elif format.lower() == "pdf":
        font = _font()
        body = ParagraphStyle("H3Body", fontName=font, fontSize=8.5, leading=12,
                              wordWrap="CJK", spaceAfter=3)
        heading = ParagraphStyle("H3Heading", parent=body, fontSize=12, leading=16,
                                 spaceBefore=12, spaceAfter=7, keepWithNext=True)
        title_style = ParagraphStyle("H3Title", parent=heading, fontSize=19, leading=24)

        def paragraph(value, style=body):
            return Paragraph(escape(_value(value)).replace("\n", "<br/>"), style)

        story = [paragraph("План полётного задания", title_style),
                 paragraph("Geoscan Mission Planner · расчётный отчёт H3"), Spacer(1, 4 * mm)]
        for title, rows in sections:
            story.append(paragraph(title, heading))
            table = Table([[paragraph(key), paragraph(value)] for key, value in rows],
                          colWidths=[54 * mm, 116 * mm], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7dedb")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef3f1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(table)

        def footer(canvas, document):
            canvas.saveState()
            canvas.setFont(font, 8)
            canvas.drawString(20 * mm, 12 * mm, "H3 · расчётная модель · не разрешение на полёт")
            canvas.drawRightString(190 * mm, 12 * mm, str(document.page))
            canvas.restoreState()

        SimpleDocTemplate(output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
                          topMargin=18 * mm, bottomMargin=20 * mm,
                          title="План полётного задания", author="Geoscan Mission Planner").build(
                              story, onFirstPage=footer, onLaterPages=footer)
    else:
        raise ValueError("Supported report formats are pdf and docx")
    return output.getvalue()
