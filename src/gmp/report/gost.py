"""ГОСТ 7.32-2017 отчёт о НИР + протокол испытаний ГОСТ 19.301 / 34.603."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

TITLE = "Интеллектуальный сервис планирования и распределения групповых авиаработ БВС"
ORG = "Хакатон Geoscan / ЛЦТ 2026"
GOST = "ГОСТ 7.32-2017, ГОСТ 19.301-79, ГОСТ 34.603-92"


def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="GostTitle", fontName="Times-Bold", fontSize=16, leading=20, alignment=1, spaceAfter=12))
    styles.add(ParagraphStyle(name="GostH", fontName="Times-Bold", fontSize=13, leading=16, spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="GostBody", fontName="Times-Roman", fontSize=11, leading=14, alignment=4, spaceAfter=6))
    styles.add(ParagraphStyle(name="GostSmall", fontName="Times-Roman", fontSize=9, leading=12, spaceAfter=4))
    return styles


def _table(data, col_widths=None):
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.9, 0.9, 0.9)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return t


def build_story(results: dict[str, Any]) -> list:
    st = _styles()
    story = []
    story.append(Paragraph("МИНИСТЕРСТВО / ОРГАНИЗАЦИЯ-ИСПОЛНИТЕЛЬ", st["GostSmall"]))
    story.append(Paragraph(ORG, st["GostTitle"]))
    story.append(Paragraph("ОТЧЁТ О НАУЧНО-ИССЛЕДОВАТЕЛЬСКОЙ РАБОТЕ", st["GostTitle"]))
    story.append(Paragraph(TITLE, st["GostH"]))
    story.append(Paragraph(f"Документ выполнен в структуре {GOST}.", st["GostBody"]))
    story.append(Paragraph(f"Дата: {results.get('generated_at', datetime.now().isoformat())}", st["GostSmall"]))
    story.append(Spacer(1, 8))

    story.append(Paragraph("1. Реферат", st["GostH"]))
    summary = results.get("summary", {})
    story.append(
        Paragraph(
            "Объект исследования — сервис предполётного планирования групповых авиаработ БВС "
            "(mission orchestration / preflight planning). Цель работы — автоматически строить "
            "полный, ресурсно допустимый и пространственно-временной безопасный план съёмки "
            "для разнородного парка до 10 физических БВС на площади порядка 100 км², а при "
            "невыполнимости исходной постановки — диагностировать причину и предложить "
            "проверяемое изменение условий. Метод — синтетический customer-conformance набор "
            "S00–S11 (seed 20260918), независимый Safety Validator, классические VRPTW и "
            f"MovingAI MAPF. Пройдено тестов: {summary.get('passed', 0)} из {summary.get('total', 0)} "
            f"(доля {summary.get('pass_rate', 0):.1f} %).",
            st["GostBody"],
        )
    )

    story.append(Paragraph("2. Введение и постановка", st["GostH"]))
    story.append(
        Paragraph(
            "Заказчик требует веб-сервис, который по подготовленной GIS-сцене, парку БВС и "
            "требованиям к съёмке (RGB, multispectral, IR, LiDAR, geophysical) формирует "
            "маршруты, расписание вылетов, оценку ресурса и Safety Certificate. Два независимых "
            "критерия оптимизации: minimum completion time (makespan) и minimum total flight. "
            "Сервис не заменяет штатный автопилот.",
            st["GostBody"],
        )
    )

    story.append(Paragraph("3. Объект и средства испытаний", st["GostH"]))
    story.append(
        Paragraph(
            "Испытаниям подвергнут пакет gmp (Geoscan Mission Planner): Coverage Engine, "
            "global optimizer (baselines / CP-SAT / LNS), multi-sortie scheduler, 4D deconfliction, "
            "независимый Safety Validator, Recommendation Engine, REST API и веб-клиент MapLibre. "
            "Внешняя СК — WGS-84 / EPSG:4326, внутренние расчёты — метрическая UTM.",
            st["GostBody"],
        )
    )

    story.append(Paragraph("4. Программа и методика испытаний (ГОСТ 19.301 / 34.603)", st["GostH"]))
    story.append(
        Paragraph(
            "Приняты пять уровней: (1) модульные и property-based тесты геометрических инвариантов; "
            "(2) Synthetic Customer-Conformance Suite S00–S11; (3) VRPTW Solomon; "
            "(4) MovingAI MAPF; (5) экспорт KML/GeoJSON. Ворота CI: A parser, B coverage, "
            "C hard safety, D resource, E reserve landing, F infeasibility, G regression, H benchmark. "
            "Оракул сценария — expected_assertions.json; готовый маршрут не навязывается solver-у.",
            st["GostBody"],
        )
    )

    story.append(Paragraph("5. Результаты customer-conformance", st["GostH"]))
    rows = [["Сценарий", "Статус плана", "Покрытие, %", "Makespan, мин", "Вылеты", "Сертификат", "Оракул"]]
    for row in results.get("conformance", []):
        m = row.get("metrics") or {}
        rows.append(
            [
                row.get("scenario", ""),
                str(row.get("status", "")),
                str(m.get("coverage_percent", "")),
                str(m.get("makespan_min", "")),
                str(m.get("sortie_count", "")),
                "да" if m.get("certificate") else "нет",
                "PASS" if row.get("passed") else "FAIL",
            ]
        )
    if len(rows) > 1:
        story.append(_table(rows, [95, 55, 50, 55, 40, 50, 40]))
    else:
        story.append(Paragraph("Таблица будет заполнена после прогона test suite.", st["GostBody"]))

    story.append(Paragraph("6. Ворота безопасности (Gate C–F)", st["GostH"]))
    gates = results.get("gates") or {}
    grow = [["Ворота", "Содержание", "Результат"]]
    for gid, g in gates.items():
        grow.append([gid, g.get("title", ""), "PASS" if g.get("ok") else "FAIL"])
    if len(grow) > 1:
        story.append(_table(grow, [50, 320, 50]))

    story.append(Paragraph("7. VRPTW и MAPF", st["GostH"]))
    story.append(
        Paragraph(
            "VRPTW (Solomon) проверяет ядро маршрутизации с временными окнами и ёмкостью, "
            "но не является UAV-датасетом. MovingAI MAPF проверяет подсистему разведения "
            "конфликтов на сетке. Оба нужны как внешние регрессии, не заменяющие Safety Validator.",
            st["GostBody"],
        )
    )
    brow = [["Бенчмарк", "Экземпляр", "Результат", "Метрика"]]
    for b in results.get("benchmarks", []):
        brow.append([b.get("suite", ""), b.get("name", ""), "PASS" if b.get("ok") else "FAIL", str(b.get("metric", ""))])
    if len(brow) > 1:
        story.append(_table(brow, [80, 140, 50, 150]))

    story.append(Paragraph("8. Выводы, пригодные для презентации", st["GostH"]))
    for bullet in results.get("slides", []):
        story.append(Paragraph(f"— {bullet}", st["GostBody"]))
    story.append(
        Paragraph(
            "Система считается готовой не когда она «рисует маршруты», а когда воспроизводимо "
            "доказывает на внешних benchmark-ах и customer-conformance suite, что строит полные, "
            "ресурсно допустимые, 4D-безопасные планы, а при невыполнимости корректно диагностирует "
            "причину и предлагает проверяемое изменение условий.",
            st["GostBody"],
        )
    )
    story.append(Paragraph("9. Приложения", st["GostH"]))
    story.append(
        Paragraph(
            "П.1 OpenAPI /health /api/v1. П.2 Safety Certificate JSON. П.3 Экспорт KML/GeoJSON. "
            "П.4 Каталог БВС ≥10 моделей. П.5 Seed 20260918, MANIFEST.json.",
            st["GostBody"],
        )
    )
    return story


def write_pdf(path: Path, results: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=15 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=TITLE,
        author=ORG,
    )
    doc.build(build_story(results))
    return path


def write_docx(path: Path, results: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.5)
    p = doc.add_paragraph(ORG)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h = doc.add_heading("ОТЧЁТ О НИР", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(f"Нормативная база: {GOST}.")
    doc.add_paragraph(f"Дата формирования: {results.get('generated_at', '')}.")
    summary = results.get("summary") or {}
    doc.add_heading("Реферат", level=1)
    doc.add_paragraph(
        f"Пройдено {summary.get('passed', 0)} из {summary.get('total', 0)} "
        f"зарегистрированных испытаний (доля {summary.get('pass_rate', 0):.1f} %)."
    )
    doc.add_heading("Результаты conformance", level=1)
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "Сценарий"
    hdr[1].text = "Статус"
    hdr[2].text = "Покрытие %"
    hdr[3].text = "Makespan"
    hdr[4].text = "Оракул"
    for row in results.get("conformance", []):
        cells = table.add_row().cells
        m = row.get("metrics") or {}
        cells[0].text = str(row.get("scenario", ""))
        cells[1].text = str(row.get("status", ""))
        cells[2].text = str(m.get("coverage_percent", ""))
        cells[3].text = str(m.get("makespan_min", ""))
        cells[4].text = "PASS" if row.get("passed") else "FAIL"
    doc.add_heading("Тезисы для презентации", level=1)
    for bullet in results.get("slides", []):
        doc.add_paragraph(bullet, style="List Bullet")
    doc.save(str(path))
    return path


def default_slides(results: dict[str, Any]) -> list[str]:
    s = results.get("summary") or {}
    slides = [
        "Задача: предполётное планирование группы разнородных БВС, не автопилот.",
        "Вход: GeoJSON/KML/GeoTIFF, 5 типов съёмки, NFZ, DEM, ветер, площадки.",
        "Выход: маршруты, расписание, ресурс, Safety Certificate, KML/GeoJSON.",
        "Два независимых objective: makespan и total flight.",
        f"Испытания: {s.get('passed', 0)}/{s.get('total', 0)} PASS, seed 20260918.",
        "Невыполнимость не маскируется: диагноз + Top-K рекомендации с re-plan.",
        "4D deconfliction: shift / hold / высота / reorder / LNS / sequentialise.",
        "Веб-служба: FastAPI + MapLibre, внешние порты 8080 и 80.",
    ]
    return slides


def write_reports(results: dict[str, Any], directory: str | Path) -> dict[str, str]:
    d = Path(directory)
    d.mkdir(parents=True, exist_ok=True)
    results = dict(results)
    results.setdefault("generated_at", datetime.now().isoformat())
    results.setdefault("slides", default_slides(results))
    (d / "test_results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    pdf = write_pdf(d / "GOST_732_test_report.pdf", results)
    docx = write_docx(d / "GOST_732_test_report.docx", results)
    return {"pdf": str(pdf), "docx": str(docx), "json": str(d / "test_results.json")}
