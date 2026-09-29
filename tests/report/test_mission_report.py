import io
import zipfile

import pytest
from docx import Document

from gmp.report.mission import generate_report_bytes


def sample():
    result = {"scene_id": "Тест МГУ", "objective": "makespan", "status": "INFEASIBLE", "sorties": [],
              "diagnosis": {"proof": {"type": "wind_excludes_all", "job_id": "RGB"}}}
    validation = {"passed": True, "status": "INFEASIBLE", "certificate": None,
                  "violations": [], "metrics": {"proof_verified": True}, "validator_version": "test"}
    metadata = {"mode": "live", "description": "<МГУ> & проверка", "metric_crs": "EPSG:32637",
                "terrain": {"source_url": "https://example.org/terrain", "attribution": "Источник высот"}}
    return result, validation, metadata


def test_pdf_has_embedded_cyrillic_font_and_real_metadata():
    data = generate_report_bytes("pdf", *sample())
    assert data.startswith(b"%PDF")
    assert b"DejaVuSans" in data
    assert len(data) > 10000


def test_docx_is_editable_and_does_not_claim_safe_or_gost():
    data = generate_report_bytes("docx", *sample())
    document = Document(io.BytesIO(data))
    text = "\n".join(c.text for table in document.tables for row in table.rows for c in row.cells)
    assert "Тест МГУ" in text and "<МГУ> & проверка" in text
    assert "Не выдан" in text and "INFEASIBLE" in text
    assert "Источник высот" in text
    assert "ГОСТ" not in text
    assert zipfile.is_zipfile(io.BytesIO(data))


def test_unknown_format_is_rejected():
    with pytest.raises(ValueError):
        generate_report_bytes("html", *sample())


def test_stale_certificate_is_not_represented_as_safe():
    result, validation, metadata = sample()
    result["status"] = "SAFE"
    validation.update(status="SAFE", certificate={"result_sha256": "old", "input_sha256": "old"})
    document = Document(io.BytesIO(generate_report_bytes("docx", result, validation, metadata)))
    text = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert "CERTIFICATE_MISMATCH" in text and "UNSAFE" in text and "Не выдан" in text


def test_auto_dem_sources_license_and_fixture_mode_are_exported():
    result, validation, metadata = sample()
    metadata["mode"] = "fixture"
    metadata["terrain"] = {"sources": [{"url": "https://example.org/dsm-tile.tif", "sha256": "ab" * 32,
                                       "retrieved_at": "2026-09-21T00:00:00Z"}],
                           "license_url": "https://example.org/license"}
    document = Document(io.BytesIO(generate_report_bytes("docx", result, validation, metadata)))
    text = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert "https://example.org/dsm-tile.tif" in text and "ab" * 32 in text
    assert "https://example.org/license" in text and "Готовый эталон" in text
    assert "финальное снятие данных" in text
